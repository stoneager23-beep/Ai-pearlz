"""Helpers that write small knowledge-base files of every supported format for tests."""

import json
from pathlib import Path

import docx
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A5
from reportlab.pdfgen import canvas


def write_pdf(path: Path, pages: list[list[str]], header: str = "Acme Manual | Rev A") -> None:
    """A PDF whose every page carries the same header and a "Page N" footer."""
    pdf = canvas.Canvas(str(path), pagesize=A5)
    for number, lines in enumerate(pages, start=1):
        pdf.drawString(30, 560, header)
        y = 520
        for line in lines:
            pdf.drawString(30, y, line)
            y -= 16
        pdf.drawString(30, 20, f"Page {number}")
        pdf.showPage()
    pdf.save()


def encrypt_pdf(source: Path, target: Path, password: str = "secret") -> None:
    writer = PdfWriter()
    for page in PdfReader(source).pages:
        writer.add_page(page)
    writer.encrypt(user_password=password)
    with target.open("wb") as handle:
        writer.write(handle)


def write_docx(path: Path, title: str, sections: dict[str, list[str]]) -> None:
    """A DOCX with a Title, then one Heading 1 per section followed by its paragraphs."""
    document = docx.Document()
    document.add_heading(title, level=0)
    for heading, paragraphs in sections.items():
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.save(path)


def write_catalog(kb_dir: Path, entries: dict[str, str]) -> None:
    """`entries` maps file name to doc_type; the other catalogue fields get fixed values."""
    catalog = {
        name: {
            "title": f"Title of {name}",
            "doc_type": doc_type,
            "product_lines": ["all"],
            "version": "1.0",
            "last_updated": "2026-01-01",
        }
        for name, doc_type in entries.items()
    }
    (kb_dir / "_catalog.json").write_text(json.dumps(catalog), encoding="utf-8")


PRICING_CSV = (
    "sku,product_name,product_line,category,description,list_price_pkr,promo_price_pkr,"
    "availability,compatible_with,notes\n"
    "F5-PF1,Sediment pre-filter PF-1,AquaPearl 500 Pro,filter,Stage 1 cartridge,2400,1990,"
    "in_stock,AquaPearl 500 Pro,\n"
    "AP250,AquaPearl 250,AquaPearl 250,purifier,Old model,,,discontinued,,Discontinued in 2024\n"
)


def build_mixed_knowledge_base(kb_dir: Path) -> None:
    """One file of each supported format, all with catalogue entries."""
    kb_dir.mkdir(parents=True, exist_ok=True)
    write_pdf(
        kb_dir / "manual.pdf",
        [
            ["Acme Purifier", "User manual cover"],
            ["1. Safety", "Use cold water only.", "2. Cleaning", "Wipe weekly."],
            ["Descale the nozzle every 3 months."],
        ],
    )
    write_docx(
        kb_dir / "agreement.docx",
        "Service Agreement",
        {"1. Plans": ["Basic plan costs PKR 6,500."], "2. Claims": ["Register a claim."]},
    )
    (kb_dir / "policy.md").write_text(
        "# Policy\n\nIntro text.\n\n## 1. Window\n\nReturn within **14 days**.\n\n"
        "## 2. Empty heading\n\n### 2.1 Detail\n\nDetail text.\n",
        encoding="utf-8",
    )
    (kb_dir / "notes.txt").write_text(
        "RELEASE NOTES\nIntro.\n\nVERSION 2.0 - released 2026\n- Change A\n", encoding="utf-8"
    )
    (kb_dir / "prices.csv").write_text(PRICING_CSV, encoding="utf-8")
    write_catalog(
        kb_dir,
        {
            "manual.pdf": "manual",
            "agreement.docx": "policy",
            "policy.md": "policy",
            "notes.txt": "release_notes",
            "prices.csv": "pricing",
        },
    )
