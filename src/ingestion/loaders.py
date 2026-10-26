"""Load the knowledge-base folder into LangChain Documents, one per structural unit (Task 02).

Files are dispatched to a loader by extension:

    .pdf   PyPDFLoader, then split at numbered headings   -> one Document per page section
    .docx  DocxSectionLoader (python-docx)                 -> one Document per heading section
    .md    TextLoader + MarkdownHeaderTextSplitter         -> one Document per heading section
    .txt   TextLoader, then split at capitalised headings  -> one Document per section
    .csv   CSVLoader                                       -> one Document per row, as a sentence

Loading by structural unit means every Document already knows where it came from (a page and
section, a section, or a row and SKU), which is what citations need later.

Every Document carries the metadata contract from manual Sec 5: `source`, `doc_type`,
`product_lines`, `version`, `last_updated`, `doc_title`, plus its location. Document-level
values come from `_catalog.json`; metadata the loaders provide (such as `page`) is kept.

Errors are contained per file: an unreadable, password-protected or empty file is logged and
skipped, and the rest of the folder still loads. Files whose names start with "_" or "." are
never loaded (the catalogue and the answer key of planted gaps live in such files).
"""

import csv
import hashlib
import json
import logging
import zipfile
from collections.abc import Callable, Collection, Iterator
from dataclasses import dataclass, field
from datetime import date
from enum import StrEnum
from pathlib import Path
from typing import Literal

import docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from pypdf.errors import FileNotDecryptedError, PdfReadError

from src.ingestion.text_cleaning import (
    collapse_whitespace,
    find_repeated_lines,
    has_body_text,
    remove_lines,
    split_numbered_sections,
    split_plain_text_sections,
    strip_markdown_emphasis,
    strip_markdown_heading_marks,
)

CATALOG_FILE_NAME = "_catalog.json"
IGNORED_NAME_PREFIXES = ("_", ".")

DocType = Literal["manual", "guide", "policy", "faq", "pricing", "release_notes"]

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """Raised when the knowledge base as a whole cannot be loaded (not for a single bad file)."""


class CatalogEntry(BaseModel):
    """Document-level metadata for one file, as written in `_catalog.json`."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    title: str = Field(min_length=1)
    doc_type: DocType
    product_lines: list[str] = Field(min_length=1)
    version: str = Field(min_length=1)
    last_updated: date


class FileStatus(StrEnum):
    LOADED = "loaded"
    SKIPPED = "skipped"


@dataclass(frozen=True)
class FileReport:
    """What happened to one file, for the ingestion report."""

    file_name: str
    file_type: str
    status: FileStatus
    reason: str = ""
    characters: int = 0
    units: int = 0
    pages: int = 0
    metadata_keys: tuple[str, ...] = ()
    stripped_lines: tuple[str, ...] = ()


@dataclass
class IngestionResult:
    documents: list[Document] = field(default_factory=list)
    reports: list[FileReport] = field(default_factory=list)

    @property
    def loaded_count(self) -> int:
        return sum(report.status is FileStatus.LOADED for report in self.reports)

    @property
    def skipped_count(self) -> int:
        return sum(report.status is FileStatus.SKIPPED for report in self.reports)


@dataclass
class _LoadedFile:
    documents: list[Document]
    pages: int = 0
    stripped_lines: tuple[str, ...] = ()


# --- Format loaders -----------------------------------------------------------


def _load_pdf(path: Path, entry: CatalogEntry) -> _LoadedFile:
    page_documents = PyPDFLoader(str(path)).load()
    page_texts = [collapse_whitespace(page.page_content) for page in page_documents]
    repeated = find_repeated_lines(page_texts)

    documents: list[Document] = []
    # Text before the first numbered heading (the cover page) is labelled with the title.
    current_heading = entry.title
    for page, text in zip(page_documents, page_texts, strict=True):
        for section in split_numbered_sections(remove_lines(text, repeated), current_heading):
            metadata = {
                **page.metadata,
                "page_number": page.metadata["page"] + 1,
                "section": section.heading,
            }
            documents.append(Document(page_content=section.text, metadata=metadata))
            current_heading = section.heading
    return _LoadedFile(documents, pages=len(page_documents), stripped_lines=tuple(sorted(repeated)))


class DocxSectionLoader(BaseLoader):
    """Load a .docx file as one Document per heading section; tables become " | " rows.

    LangChain's Docx2txtLoader flattens the file to plain text and loses the headings, so it
    cannot say which section a sentence came from. This loader walks paragraphs and tables in
    document order and starts a new Document at every Word heading.
    """

    def __init__(self, file_path: Path) -> None:
        self.file_path = file_path

    def lazy_load(self) -> Iterator[Document]:
        word_document = docx.Document(str(self.file_path))
        title = ""
        headings: dict[int, str] = {}
        lines: list[str] = []
        for block in word_document.iter_inner_content():
            if isinstance(block, DocxTable):
                lines.extend(
                    " | ".join(cell.text.strip() for cell in row.cells) for row in block.rows
                )
                continue
            level = _docx_heading_level(block)
            if level is None:
                if block.text.strip():
                    lines.append(block.text)
                continue
            yield from self._section_document(title, headings, lines)
            lines = [block.text]
            if level == 0:
                title = block.text
            else:
                headings = {depth: text for depth, text in headings.items() if depth < level}
                headings[level] = block.text
        yield from self._section_document(title, headings, lines)

    def _section_document(
        self, title: str, headings: dict[int, str], lines: list[str]
    ) -> Iterator[Document]:
        heading = headings[max(headings)] if headings else title
        text = "\n".join(lines)
        if has_body_text(text, heading):
            path = " > ".join(headings[depth] for depth in sorted(headings)) or title
            metadata = {"source": str(self.file_path), "section": heading, "section_path": path}
            yield Document(page_content=text, metadata=metadata)


def _docx_heading_level(paragraph: DocxParagraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name == "Title":
        return 0
    if style_name.startswith("Heading "):
        return int(style_name.removeprefix("Heading "))
    return None


def _load_docx(path: Path, entry: CatalogEntry) -> _LoadedFile:
    documents = DocxSectionLoader(path).load()
    for document in documents:
        document.page_content = collapse_whitespace(document.page_content)
    return _LoadedFile(documents)


_MARKDOWN_HEADINGS = [("#", "h1"), ("##", "h2"), ("###", "h3")]


def _load_markdown(path: Path, entry: CatalogEntry) -> _LoadedFile:
    raw = TextLoader(str(path), encoding="utf-8").load()[0]
    splitter = MarkdownHeaderTextSplitter(_MARKDOWN_HEADINGS, strip_headers=False)
    documents = []
    for section in splitter.split_text(strip_markdown_emphasis(raw.page_content)):
        levels = [section.metadata[key] for key in ("h2", "h3") if key in section.metadata]
        heading = levels[-1] if levels else section.metadata.get("h1", "")
        text = collapse_whitespace(strip_markdown_heading_marks(section.page_content))
        if has_body_text(text, heading):
            metadata = {
                **raw.metadata,
                "section": heading,
                "section_path": " > ".join(levels) or heading,
            }
            documents.append(Document(page_content=text, metadata=metadata))
    return _LoadedFile(documents)


def _load_plain_text(path: Path, entry: CatalogEntry) -> _LoadedFile:
    raw = TextLoader(str(path), encoding="utf-8").load()[0]
    documents = [
        Document(page_content=section.text, metadata={**raw.metadata, "section": section.heading})
        for section in split_plain_text_sections(collapse_whitespace(raw.page_content))
    ]
    return _LoadedFile(documents)


def _load_csv(path: Path, entry: CatalogEntry) -> _LoadedFile:
    with path.open(encoding="utf-8", newline="") as handle:
        columns = next(csv.reader(handle))
    rows = CSVLoader(str(path), encoding="utf-8", metadata_columns=columns).load()
    describe = describe_price_row if entry.doc_type == "pricing" else _describe_generic_row

    documents = []
    for row in rows:
        values = {column: (row.metadata.get(column) or "").strip() for column in columns}
        metadata = {"source": row.metadata["source"], "row": row.metadata["row"]}
        if entry.doc_type == "pricing":
            metadata |= {
                "sku": values["sku"],
                "product_name": values["product_name"],
                "category": values["category"],
                "availability": values["availability"],
                "row_product_line": values["product_line"],
            }
        documents.append(Document(page_content=describe(values), metadata=metadata))
    return _LoadedFile(documents)


_CATEGORY_WORDS = {
    "purifier": "water purifier",
    "air_purifier": "air purifier",
    "filter": "replacement filter",
    "bundle": "bundle",
    "accessory": "accessory",
    "service": "service",
}
_AVAILABILITY_WORDS = {
    "in_stock": "in stock",
    "low_stock": "low stock",
    "out_of_stock": "out of stock",
    "discontinued": "discontinued, no longer sold",
    "available": "available to book",
}


def _pkr(amount: str) -> str:
    return f"PKR {int(amount):,}"


def describe_price_row(values: dict[str, str]) -> str:
    """Rewrite one pricing row as a self-describing sentence (manual mistake #13).

    A chunk then always carries the product name and SKU next to its price, so a price can
    never be retrieved without knowing what it is the price of.
    """
    category = _CATEGORY_WORDS.get(values["category"], values["category"])
    line = values["product_line"]
    for_whom = "all product lines" if line == "all" else f"the {line}"
    parts = [
        f"{values['product_name']} (SKU {values['sku']}), {category} for {for_whom}.",
        f"{values['description']}.",
    ]
    if values["compatible_with"]:
        parts.append(f"Compatible with: {values['compatible_with']}.")
    if values["list_price_pkr"]:
        parts.append(f"List price: {_pkr(values['list_price_pkr'])}.")
    else:
        parts.append("No list price: this item is no longer sold.")
    if values["promo_price_pkr"]:
        parts.append(f"Current promotional price: {_pkr(values['promo_price_pkr'])}.")
    else:
        parts.append("No promotional price is currently offered.")
    availability = _AVAILABILITY_WORDS.get(values["availability"], values["availability"])
    parts.append(f"Availability: {availability}.")
    if values["notes"]:
        parts.append(f"Note: {values['notes']}.")
    return " ".join(parts)


def _describe_generic_row(values: dict[str, str]) -> str:
    return "; ".join(f"{column}: {value}" for column, value in values.items() if value) + "."


_Loader = Callable[[Path, CatalogEntry], _LoadedFile]

LOADERS: dict[str, _Loader] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_markdown,
    ".txt": _load_plain_text,
    ".csv": _load_csv,
}


# --- Folder ingestion ---------------------------------------------------------


def read_catalog(kb_dir: Path) -> dict[str, CatalogEntry]:
    """Read and validate `_catalog.json`. A missing or invalid catalogue stops ingestion."""
    catalog_path = kb_dir / CATALOG_FILE_NAME
    if not catalog_path.is_file():
        raise IngestionError(f"{catalog_path} not found; every document needs a catalogue entry.")
    try:
        raw = json.loads(catalog_path.read_text(encoding="utf-8"))
        return {name: CatalogEntry.model_validate(entry) for name, entry in raw.items()}
    except (json.JSONDecodeError, ValidationError) as error:
        raise IngestionError(f"{catalog_path.name} is invalid: {error}") from error


def _skip_reason(error: Exception) -> str:
    """A short, human description of why a file could not be read."""
    if isinstance(error, FileNotDecryptedError):
        return "password-protected PDF"
    if isinstance(error, PdfReadError):
        return f"unreadable or corrupt PDF ({error})"
    if isinstance(error, zipfile.BadZipFile):
        return "corrupt DOCX (not a valid Word file)"
    if isinstance(error, UnicodeDecodeError):
        return "not UTF-8 text"
    return f"{type(error).__name__}: {error}"


def _apply_metadata_contract(
    document: Document, relative_name: str, entry: CatalogEntry
) -> Document:
    metadata = dict(document.metadata)
    metadata["file_path"] = metadata.pop("source", relative_name)
    metadata |= {
        "source": relative_name,
        "doc_title": entry.title,
        "doc_type": entry.doc_type,
        "product_lines": metadata.pop("row_product_line", None) or entry.product_lines,
        "version": entry.version,
        "last_updated": entry.last_updated.isoformat(),
    }
    if isinstance(metadata["product_lines"], str):
        metadata["product_lines"] = [metadata["product_lines"]]
    return Document(page_content=document.page_content, metadata=metadata)


def _content_hash(documents: list[Document]) -> str:
    joined = "\n".join(document.page_content for document in documents)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


def _is_ignored(path: Path, kb_dir: Path) -> bool:
    return any(part.startswith(IGNORED_NAME_PREFIXES) for part in path.relative_to(kb_dir).parts)


def load_knowledge_base(kb_dir: Path, only: Collection[str] | None = None) -> IngestionResult:
    """Load every supported file under `kb_dir` (recursively) into Documents.

    `only`, if given, restricts loading to those file names (an incremental index update
    re-reads just the files that changed). One bad file never stops the run: it is logged
    with a warning and reported as skipped.
    """
    if not kb_dir.is_dir():
        raise IngestionError(f"Knowledge-base folder {kb_dir} does not exist.")
    catalog = read_catalog(kb_dir)
    result = IngestionResult()
    seen_hashes: dict[str, str] = {}

    for path in sorted(kb_dir.rglob("*")):
        if not path.is_file() or _is_ignored(path, kb_dir):
            continue
        name = path.relative_to(kb_dir).as_posix()
        if only is not None and name not in only:
            continue
        file_type = path.suffix.lower().lstrip(".")

        def skip(reason: str, name: str = name, file_type: str = file_type) -> None:
            logger.warning("Skipping %s: %s", name, reason)
            result.reports.append(FileReport(name, file_type, FileStatus.SKIPPED, reason))

        loader = LOADERS.get(path.suffix.lower())
        if loader is None:
            skip(f"unsupported file type .{file_type}")
            continue
        entry = catalog.get(name)
        if entry is None:
            skip(f"no entry in {CATALOG_FILE_NAME}")
            continue
        try:
            loaded = loader(path, entry)
        # Per-file isolation: third-party parsers raise many exception types for bad files,
        # and one bad file must never stop the whole knowledge base from loading.
        except Exception as error:
            skip(_skip_reason(error))
            continue

        documents = [doc for doc in loaded.documents if doc.page_content.strip()]
        if not documents:
            skip("no text could be extracted")
            continue
        content_hash = _content_hash(documents)
        if content_hash in seen_hashes:
            skip(f"duplicate of {seen_hashes[content_hash]}")
            continue
        seen_hashes[content_hash] = name

        documents = [_apply_metadata_contract(doc, name, entry) for doc in documents]
        result.documents.extend(documents)
        result.reports.append(
            FileReport(
                file_name=name,
                file_type=file_type,
                status=FileStatus.LOADED,
                characters=sum(len(doc.page_content) for doc in documents),
                units=len(documents),
                pages=loaded.pages,
                metadata_keys=tuple(sorted({key for doc in documents for key in doc.metadata})),
                stripped_lines=loaded.stripped_lines,
            )
        )
        logger.info(
            "Loaded %s: %d units, %d characters",
            name,
            len(documents),
            result.reports[-1].characters,
        )

    logger.info(
        "Ingestion finished: %d files found, %d loaded, %d skipped",
        len(result.reports),
        result.loaded_count,
        result.skipped_count,
    )
    return result
