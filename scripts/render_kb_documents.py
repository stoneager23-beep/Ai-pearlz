"""Render the PDF and DOCX knowledge-base documents from their Markdown sources.

The Pearlz manual, installation guide, service agreement and privacy policy are authored as
plain Markdown in `data/kb_sources/` so they are easy to review and edit. This script turns
each one into the format a real company would publish it in (PDF or DOCX) and writes it to
`data/knowledge_base/`, where ingestion picks it up. The sources themselves are never ingested.

Each source starts with a small front-matter block:

    ---
    output: product_manual_aquapearl500.pdf
    header: AquaPearl 500 Pro User Manual  |  Revision C
    ---

Supported Markdown: `#`/`##`/`###` headings, paragraphs, `- ` bullets, `1. ` numbered items,
pipe tables, `> ` callouts, `**bold**`, and `<!-- pagebreak -->`.

Usage:  python -m scripts.render_kb_documents
"""

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape

import reportlab
from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.fonts import addMapping
from reportlab.lib.pagesizes import A5
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = PROJECT_ROOT / "data" / "kb_sources"
OUTPUT_DIR = PROJECT_ROOT / "data" / "knowledge_base"

PAGE_BREAK_MARKER = "<!-- pagebreak -->"
_PAGE_MARGIN = 14 * mm
_HEADER_FONT_SIZE = 7
_FOOTER_TEXT = "Pearlz Home Systems (Pvt.) Ltd.  |  pearlzhome.example"
# A narrow column (such as "Code") still gets at least this share of the widest column's width.
_MIN_COLUMN_SHARE = 0.35
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMBERED_ITEM = re.compile(r"^\d+\.\s+")

logger = logging.getLogger(__name__)


BlockKind = Literal[
    "title",
    "heading1",
    "heading2",
    "paragraph",
    "bullets",
    "numbered",
    "table",
    "callout",
    "pagebreak",
]


@dataclass
class Block:
    """One structural element of a source document."""

    kind: BlockKind
    text: str = ""
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class SourceDocument:
    output_name: str
    header: str
    blocks: list[Block]


def parse_source(path: Path) -> SourceDocument:
    """Split a Markdown source into its front matter and a list of blocks."""
    lines = path.read_text(encoding="utf-8").splitlines()
    front_matter, body = _split_front_matter(lines)
    return SourceDocument(
        output_name=front_matter["output"],
        header=front_matter["header"],
        blocks=_parse_blocks(body),
    )


def _split_front_matter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    if not lines or lines[0].strip() != "---":
        raise ValueError("source must start with a '---' front-matter block")
    closing = lines.index("---", 1)
    values = {}
    for line in lines[1:closing]:
        key, _, value = line.partition(":")
        values[key.strip()] = value.strip()
    return values, lines[closing + 1 :]


def _parse_blocks(lines: list[str]) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(Block("paragraph", text=" ".join(paragraph)))
            paragraph.clear()

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue
        if line == PAGE_BREAK_MARKER:
            flush_paragraph()
            blocks.append(Block("pagebreak"))
        elif line.startswith("### "):
            flush_paragraph()
            blocks.append(Block("heading2", text=line[4:]))
        elif line.startswith("## "):
            flush_paragraph()
            blocks.append(Block("heading1", text=line[3:]))
        elif line.startswith("# "):
            flush_paragraph()
            blocks.append(Block("title", text=line[2:]))
        elif line.startswith("|"):
            flush_paragraph()
            _append_table_row(blocks, line)
        elif line.startswith("- "):
            flush_paragraph()
            _append_list_item(blocks, "bullets", line[2:])
        elif _NUMBERED_ITEM.match(line):
            flush_paragraph()
            _append_list_item(blocks, "numbered", _NUMBERED_ITEM.sub("", line))
        elif line.startswith("> "):
            flush_paragraph()
            blocks.append(Block("callout", text=line[2:]))
        else:
            paragraph.append(line)
    flush_paragraph()
    return blocks


def _append_table_row(blocks: list[Block], line: str) -> None:
    cells = [cell.strip() for cell in line.strip("|").split("|")]
    if all(set(cell) <= set("-: ") for cell in cells):
        return  # the |---|---| separator under the header row
    if not blocks or blocks[-1].kind != "table":
        blocks.append(Block("table"))
    blocks[-1].rows.append(cells)


def _append_list_item(blocks: list[Block], kind: str, text: str) -> None:
    if not blocks or blocks[-1].kind != kind:
        blocks.append(Block(kind))
    blocks[-1].items.append(text)


# --- PDF ---------------------------------------------------------------------

# reportlab's built-in Helvetica has no Unicode map, so PDF text extraction turns characters
# such as the degree sign into "?". The TrueType fonts bundled with reportlab embed one.
_FONT = "Vera"
_FONT_BOLD = "VeraBd"


def _register_fonts() -> None:
    font_dir = Path(reportlab.__file__).parent / "fonts"
    for name, file_name in (
        (_FONT, "Vera.ttf"),
        (_FONT_BOLD, "VeraBd.ttf"),
        ("VeraIt", "VeraIt.ttf"),
        ("VeraBI", "VeraBI.ttf"),
    ):
        pdfmetrics.registerFont(TTFont(name, str(font_dir / file_name)))
    addMapping(_FONT, 0, 0, _FONT)
    addMapping(_FONT, 1, 0, _FONT_BOLD)
    addMapping(_FONT, 0, 1, "VeraIt")
    addMapping(_FONT, 1, 1, "VeraBI")


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body", parent=base["BodyText"], fontName=_FONT, fontSize=8.5, leading=12, spaceAfter=5
    )
    return {
        "title": ParagraphStyle(
            "Title", parent=base["Title"], fontName=_FONT_BOLD, fontSize=20, leading=24
        ),
        "heading1": ParagraphStyle(
            "H1",
            parent=base["Heading1"],
            fontName=_FONT_BOLD,
            fontSize=12.5,
            leading=16,
            spaceBefore=6,
        ),
        "heading2": ParagraphStyle(
            "H2",
            parent=base["Heading2"],
            fontName=_FONT_BOLD,
            fontSize=10,
            leading=13,
            spaceBefore=4,
        ),
        "paragraph": body,
        "cell": ParagraphStyle("Cell", parent=body, fontSize=7.5, leading=9.5, spaceAfter=0),
        "callout": ParagraphStyle(
            "Callout",
            parent=body,
            backColor=colors.HexColor("#EEF3F6"),
            borderPadding=5,
            leftIndent=5,
            rightIndent=5,
            spaceBefore=4,
            spaceAfter=9,
        ),
    }


def _inline_pdf(text: str) -> str:
    return _BOLD.sub(r"<b>\1</b>", escape(text))


def _column_widths(rows: list[list[str]], available_width: float) -> list[float]:
    """Share the width in proportion to each column's average text length, with a floor."""
    column_count = max(len(row) for row in rows)
    average_lengths = [
        sum(len(row[index]) for row in rows if index < len(row)) / len(rows)
        for index in range(column_count)
    ]
    floor = max(average_lengths) * _MIN_COLUMN_SHARE
    weights = [max(length, floor) for length in average_lengths]
    return [available_width * weight / sum(weights) for weight in weights]


def _pdf_table(rows: list[list[str]], styles: dict[str, ParagraphStyle]) -> Table:
    available_width = A5[0] - 2 * _PAGE_MARGIN
    cells = [[Paragraph(_inline_pdf(cell), styles["cell"]) for cell in row] for row in rows]
    table = Table(cells, colWidths=_column_widths(rows, available_width), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE7EC")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA9B1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def _pdf_flowables(blocks: list[Block], styles: dict[str, ParagraphStyle]) -> list:
    flowables: list = []
    for block in blocks:
        if block.kind == "pagebreak":
            flowables.append(PageBreak())
        elif block.kind == "table":
            flowables.extend([_pdf_table(block.rows, styles), Spacer(1, 6)])
        elif block.kind in ("bullets", "numbered"):
            items = [
                ListItem(Paragraph(_inline_pdf(item), styles["paragraph"])) for item in block.items
            ]
            bullet_type = "bullet" if block.kind == "bullets" else "1"
            flowables.append(
                ListFlowable(items, bulletType=bullet_type, leftIndent=12, bulletFontName=_FONT)
            )
        else:
            flowables.append(Paragraph(_inline_pdf(block.text), styles[block.kind]))
    return flowables


def render_pdf(source: SourceDocument, output_path: Path) -> int:
    """Write `source` as an A5 PDF with a running header and footer. Returns the page count."""
    pages: list[int] = []

    def draw_header_and_footer(canvas, document) -> None:  # noqa: ANN001 (reportlab callback)
        canvas.saveState()
        canvas.setFont(_FONT, _HEADER_FONT_SIZE)
        top = A5[1] - _PAGE_MARGIN / 2
        canvas.drawString(_PAGE_MARGIN, top, source.header)
        canvas.drawString(_PAGE_MARGIN, _PAGE_MARGIN / 2, _FOOTER_TEXT)
        canvas.drawRightString(A5[0] - _PAGE_MARGIN, _PAGE_MARGIN / 2, f"Page {document.page}")
        canvas.restoreState()
        pages.append(document.page)

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A5,
        leftMargin=_PAGE_MARGIN,
        rightMargin=_PAGE_MARGIN,
        topMargin=_PAGE_MARGIN,
        bottomMargin=_PAGE_MARGIN,
        title=source.header,
        author="Pearlz Home Systems (Pvt.) Ltd.",
    )
    flowables = _pdf_flowables(source.blocks, _pdf_styles())
    document.build(
        flowables, onFirstPage=draw_header_and_footer, onLaterPages=draw_header_and_footer
    )
    return max(pages)


# --- DOCX --------------------------------------------------------------------


def _add_inline_runs(paragraph, text: str) -> None:  # noqa: ANN001 (python-docx paragraph)
    for index, part in enumerate(_BOLD.split(text)):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def render_docx(source: SourceDocument, output_path: Path) -> int:
    """Write `source` as a DOCX using Word's built-in heading styles. Returns the block count."""
    document = DocxDocument()
    document.styles["Normal"].font.size = Pt(10.5)
    document.core_properties.author = "Pearlz Home Systems (Pvt.) Ltd."
    document.core_properties.title = source.header
    document.sections[0].header.paragraphs[0].text = source.header

    for block in source.blocks:
        if block.kind == "pagebreak":
            document.add_page_break()
        elif block.kind == "title":
            document.add_heading(block.text, level=0)
        elif block.kind == "heading1":
            document.add_heading(block.text, level=1)
        elif block.kind == "heading2":
            document.add_heading(block.text, level=2)
        elif block.kind in ("bullets", "numbered"):
            style = "List Bullet" if block.kind == "bullets" else "List Number"
            for item in block.items:
                _add_inline_runs(document.add_paragraph(style=style), item)
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=max(len(row) for row in block.rows))
            table.style = "Table Grid"
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=False):
                    cell.text = ""
                    _add_inline_runs(cell.paragraphs[0], text)
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        else:
            paragraph = document.add_paragraph()
            if block.kind == "callout":
                paragraph.style = "Intense Quote"
            _add_inline_runs(paragraph, block.text)
    document.save(output_path)
    return len(source.blocks)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    # Fixed timestamps and document IDs, so unchanged sources give byte-identical PDFs.
    rl_config.invariant = 1
    _register_fonts()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for source_path in sorted(SOURCE_DIR.glob("*.md")):
        source = parse_source(source_path)
        output_path = OUTPUT_DIR / source.output_name
        if output_path.suffix == ".pdf":
            pages = render_pdf(source, output_path)
            logger.info("%-40s -> %s (%d pages)", source_path.name, output_path.name, pages)
        elif output_path.suffix == ".docx":
            blocks = render_docx(source, output_path)
            logger.info("%-40s -> %s (%d blocks)", source_path.name, output_path.name, blocks)
        else:
            raise ValueError(f"{source_path.name}: unsupported output type {output_path.suffix}")


if __name__ == "__main__":
    main()
